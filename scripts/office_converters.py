"""MarkItDown Office converters that replace embedded images in place."""
from __future__ import annotations

from typing import Any

from markitdown._stream_info import StreamInfo
from markitdown.converters._docx_converter import DocxConverter
from markitdown.converters._pptx_converter import PptxConverter
from markitdown.converters._xlsx_converter import XlsxConverter

from assets import AssetWriter, FragmentHtmlConverter


def _image_bytes(image_stream: Any) -> bytes:
    """Read a borrowed MarkItDown image stream and leave it at its start."""
    try:
        image_stream.seek(0)
        return image_stream.read()
    finally:
        image_stream.seek(0)


def _mime(stream_info: StreamInfo) -> str | None:
    return getattr(stream_info, "mimetype", None)


def _pptx_shape_alt(shape: Any) -> str:
    try:
        value = shape._element._nvXxPr.cNvPr.attrib.get("descr", "")
    except (AttributeError, KeyError, TypeError):
        value = ""
    return value or getattr(shape, "name", "") or ""


class AssetPptxConverter(PptxConverter):
    """Preserve MarkItDown's slide structure while exporting each picture."""

    def __init__(self, writer: AssetWriter):
        super().__init__()
        self._html_converter = FragmentHtmlConverter()
        self.writer = writer
        self._shape_refs: dict[tuple[str, int], str] = {}
        self._chart_refs: dict[str, list[str]] = {}
        self._current_source_ref = "slide:unknown/shape:unknown"
        self._current_alt = ""

    def _index_shapes(self, presentation: Any) -> None:
        self._shape_refs.clear()
        self._chart_refs.clear()

        def visit(shapes: Any, slide_number: int, partname: str) -> None:
            for shape in shapes:
                source_ref = f"slide:{slide_number}/shape:{shape.shape_id}"
                self._shape_refs[(partname, shape.shape_id)] = source_ref
                try:
                    if shape.has_chart:
                        chart_part = str(shape.chart.part.partname)
                        self._chart_refs.setdefault(chart_part, []).append(source_ref)
                except (AttributeError, KeyError, ValueError):
                    pass
                try:
                    from pptx.enum.shapes import MSO_SHAPE_TYPE
                    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                        visit(shape.shapes, slide_number, partname)
                except (AttributeError, ImportError, TypeError):
                    pass

        for slide_number, slide in enumerate(presentation.slides, 1):
            visit(slide.shapes, slide_number, str(slide.part.partname))

    def convert(self, file_stream: Any, stream_info: StreamInfo, **kwargs: Any):
        import pptx

        try:
            file_stream.seek(0)
            presentation = pptx.Presentation(file_stream)
            self._index_shapes(presentation)
        finally:
            file_stream.seek(0)
        return super().convert(file_stream, stream_info, **kwargs)

    def _convert_picture_to_markdown(self, shape: Any, **kwargs: Any) -> str:
        partname = str(shape.part.partname)
        source_ref = self._shape_refs.get(
            (partname, shape.shape_id), f"slide:unknown/shape:{shape.shape_id}"
        )
        image_blob, _image_mime, _filename = self._get_image_info(shape)
        if image_blob is None:
            raise ValueError(f"{source_ref}: embedded picture has no readable image data")
        previous_ref = self._current_source_ref
        previous_alt = self._current_alt
        self._current_source_ref = source_ref
        self._current_alt = _pptx_shape_alt(shape)
        try:
            return super()._convert_picture_to_markdown(shape, **kwargs)
        finally:
            self._current_source_ref = previous_ref
            self._current_alt = previous_alt

    def _image_to_html(self, image_stream: Any, stream_info: StreamInfo, **kwargs: Any) -> str:
        data = _image_bytes(image_stream)
        return self.writer.add_image(
            data, _mime(stream_info), self._current_source_ref, alt=self._current_alt
        )

    def _convert_chart_to_markdown(self, chart: Any) -> str:
        try:
            markdown = super()._convert_chart_to_markdown(chart)
        except Exception as exc:
            refs = self._chart_refs.get(str(chart.part.partname), [])
            where = refs[0] if refs else "slide:unknown/chart"
            raise ValueError(f"{where}: chart conversion failed: {exc}") from exc
        if not isinstance(markdown, str) or "[unsupported chart]" in markdown:
            refs = self._chart_refs.get(str(chart.part.partname), [])
            where = refs[0] if refs else "slide:unknown/chart"
            raise ValueError(f"{where}: unsupported chart type")
        return markdown.rstrip() + "\n\n"

class AssetDocxConverter(DocxConverter):
    """Keep Mammoth's image positions while materializing each DOCX image."""

    def __init__(self, writer: AssetWriter):
        super().__init__()
        self._html_converter = FragmentHtmlConverter()
        self.writer = writer
        self._image_number = 0

    def convert(self, file_stream: Any, stream_info: StreamInfo, **kwargs: Any):
        self._image_number = 0
        file_stream.seek(0)
        try:
            return super().convert(file_stream, stream_info, **kwargs)
        finally:
            file_stream.seek(0)

    def _image_to_html(self, image_stream: Any, stream_info: StreamInfo, **kwargs: Any) -> str:
        self._image_number += 1
        source_ref = f"docx:image:{self._image_number}"
        data = _image_bytes(image_stream)
        alt = getattr(stream_info, "filename", None) or ""
        return self.writer.add_image(data, _mime(stream_info), source_ref, alt=alt)


class AssetXlsxConverter(XlsxConverter):
    """Keep MarkItDown's per-sheet ordering and append images after each table."""

    def __init__(self, writer: AssetWriter):
        super().__init__()
        self._html_converter = FragmentHtmlConverter()
        self.writer = writer
        self._image_number = 0

    def convert(self, file_stream: Any, stream_info: StreamInfo, **kwargs: Any):
        self._image_number = 0
        file_stream.seek(0)
        try:
            return super().convert(file_stream, stream_info, **kwargs)
        finally:
            file_stream.seek(0)

    def _image_to_html(self, image_stream: Any, stream_info: StreamInfo, **kwargs: Any) -> str:
        self._image_number += 1
        source_ref = f"xlsx:image:{self._image_number}"
        data = _image_bytes(image_stream)
        alt = getattr(stream_info, "filename", None) or ""
        return self.writer.add_image(data, _mime(stream_info), source_ref, alt=alt)
