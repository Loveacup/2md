"""End-to-end native document/asset delivery and resource confinement."""
from io import BytesIO
import base64
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import zipfile

import pytest
from PIL import Image

RUNNER = Path(__file__).resolve().parents[1] / 'scripts' / '2md.py'


def convert(tmp_path, source, *args, out='out'):
    home = tmp_path / 'home'
    home.mkdir(exist_ok=True)
    env = dict(os.environ, HOME=str(home), USERPROFILE=str(home), JZ2MD_VENV=sys.prefix)
    env.pop('JZ2MD_MARKER_VENV', None)
    p = subprocess.run([sys.executable, str(RUNNER), 'convert', str(source), '--output-dir', str(tmp_path / out), '--vision', 'off', *args], env=env, capture_output=True, text=True, timeout=90)
    return p, tmp_path / out / source.stem


def png(color):
    b = BytesIO()
    Image.new('RGB', (80, 40), color).save(b, 'PNG')
    return b.getvalue()


def delivered(p, dest, stem):
    assert p.returncode == 0, p.stdout + p.stderr
    report = json.loads((dest / 'conversion.json').read_text())
    assert report['status'] == 'success'
    return (dest / f'{stem}.md').read_text(), report


def failed(p, dest, stem):
    assert p.returncode == 1, p.stdout + p.stderr
    assert not (dest / f'{stem}.md').exists()
    report = json.loads((dest / 'conversion.json').read_text())
    assert report['status'] == 'failed'
    assert report['errors']
    assert list(dest.rglob('*.partial.md'))
    return report


def test_office_preserves_native_structure_and_image_occurrences(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    first, second = png('red'), png('blue')
    prs = Presentation()
    for n, data in enumerate((first, second), 1):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = f'NATIVE {n}'
        pic = slide.shapes.add_picture(BytesIO(data), Inches(1), Inches(2), width=Inches(1)); pic.name = 'Image0'
        group = slide.shapes.add_group_shape()
        group.shapes.add_picture(BytesIO(first), Inches(3), Inches(2), width=Inches(1)).name = 'Image0'
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(3), Inches(1)).table
        table.cell(0,0).text = 'Alpha';table.cell(0,1).text = '17'
        table.cell(1,0).text = 'Beta';table.cell(1,1).text = '29'
        chart = CategoryChartData();chart.categories = ['Alpha','Beta'];chart.add_series('Scores',[17,29])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(2), Inches(3), Inches(2), chart)
        slide.notes_slide.notes_text_frame.text = f'NOTE-42-{n}'
    source = tmp_path / 'native.pptx';prs.save(source)
    p, dest = convert(tmp_path, source)
    md, report = delivered(p,dest,source.stem)
    for value in ('NATIVE 1','NATIVE 2','Alpha','Beta','17','29','Scores','NOTE-42-1','NOTE-42-2'):
        assert value in md
    assert md.count('<!-- Slide number:') == 2
    from markdown_it import MarkdownIt
    from bs4 import BeautifulSoup
    rendered = BeautifulSoup(MarkdownIt().enable('table').render(md), 'html.parser')
    assert [[cell.get_text() for cell in table.select('thead th')] for table in rendered.find_all('table')] == [['Category', 'Scores'], ['Alpha', '17']] * 2
    assert len(report['assets']) == 2
    refs = [ref for a in report['assets'] for ref in a['source_refs']]
    assert len(refs) == 4 and len(set(refs)) == 4
    for data in (first,second):
        name = hashlib.sha256(data).hexdigest() + '.png'
        assert (dest/'assets'/name).read_bytes() == data
    assert md.count(hashlib.sha256(first).hexdigest()) == 3


@pytest.mark.parametrize('kind',['docx','xlsx'])
def test_table_and_sheet_images_stay_with_native_text(tmp_path, kind):
    image = tmp_path/'image.png';image.write_bytes(png('green'))
    source = tmp_path/f'native.{kind}'
    if kind == 'docx':
        from docx import Document
        from docx.shared import Inches
        doc=Document();t=doc.add_table(rows=1,cols=2);t.cell(0,0).text='Alpha 17'
        t.cell(0,1).paragraphs[0].add_run().add_picture(str(image),width=Inches(1));doc.add_paragraph('Beta 29');doc.save(source)
    else:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        wb=Workbook();ws=wb.active;ws.title='First';ws.append(['Alpha',17]);ws.append(['Beta',29]);ws.add_image(XLImage(str(image)),'A5');wb.save(source)
    p,dest=convert(tmp_path,source);md,report=delivered(p,dest,source.stem)
    assert all(x in md for x in ('Alpha','17','Beta','29'))
    assert len(report['assets']) == 1
    assert (dest/report['assets'][0]['path']).read_bytes() == image.read_bytes()


def test_html_local_and_data_resources_are_real_and_deduplicated(tmp_path):
    data=png('orange');(tmp_path/'local.png').write_bytes(data)
    source=tmp_path/'native.html'
    source.write_text(f'<h1>Local</h1><img src="local.png?q=ignored#x"><img src="data:image/png;base64,{base64.b64encode(data).decode()}">')
    p,dest=convert(tmp_path,source);md,report=delivered(p,dest,source.stem)
    assert len(report['assets']) == 1
    assert len(report['assets'][0]['source_refs']) == 2
    assert md.count(hashlib.sha256(data).hexdigest()) == 2


@pytest.mark.parametrize('resource',['missing.png','https://example.invalid/picture.png','%2e%2e/escape.png','/etc/passwd'])
def test_html_missing_remote_and_escaping_resources_fail_transaction(tmp_path,resource):
    source=tmp_path/'bad.html';source.write_text(f'<h1>Before</h1><img src="{resource}">')
    p,dest=convert(tmp_path,source);failed(p,dest,source.stem)


def test_html_symlink_escape_is_refused(tmp_path):
    outside=tmp_path/'outside';outside.mkdir();(outside/'x.png').write_bytes(png('red'))
    inside=tmp_path/'inside';inside.mkdir();(inside/'x.png').symlink_to(outside/'x.png')
    source=inside/'bad.html';source.write_text('<img src="x.png">')
    p,dest=convert(tmp_path,source);failed(p,dest,source.stem)


def make_epub(path,duplicate=False):
    with zipfile.ZipFile(path,'w') as z:
        z.writestr('META-INF/container.xml','<container><rootfile full-path="Book/package.opf"/></container>')
        z.writestr('Book/package.opf','<package xmlns:dc="http://purl.org/dc/elements/1.1/"><metadata><dc:title>Book</dc:title></metadata><manifest><item id="a" href="A/chapter.xhtml"/><item id="b" href="B/chapter.xhtml"/></manifest><spine><itemref idref="a"/><itemref idref="b"/></spine></package>')
        z.writestr('Book/A/chapter.xhtml','<h1>First</h1><img src="../Images/same.png">')
        z.writestr('Book/B/chapter.xhtml','<h1>Second</h1><img src="same.png">')
        z.writestr('Book/Images/same.png',png('red'));z.writestr('Book/B/same.png',png('blue'))
        if duplicate:
            with pytest.warns(UserWarning):z.writestr('Book/B/same.png',png('red'))


def test_epub_uses_chapter_paths_not_basename(tmp_path):
    source=tmp_path/'book.epub';make_epub(source)
    p,dest=convert(tmp_path,source);md,report=delivered(p,dest,source.stem)
    assert md.index('First') < md.index(hashlib.sha256(png('red')).hexdigest()) < md.index('Second') < md.index(hashlib.sha256(png('blue')).hexdigest())
    assert len(report['assets']) == 2


def test_duplicate_epub_members_fail(tmp_path):
    source=tmp_path/'book.epub';make_epub(source,True)
    p,dest=convert(tmp_path,source);failed(p,dest,source.stem)


def test_multiframe_preserves_each_frame(tmp_path):
    source=tmp_path/'frames.tiff'
    frames=[Image.new('RGB',(80,40),c) for c in ('red','blue')]
    frames[0].save(source,save_all=True,append_images=frames[1:])
    p,dest=convert(tmp_path,source);md,report=delivered(p,dest,source.stem)
    assert len(report['assets']) == 2
    colors=[]
    for asset in report['assets']:
        with Image.open(dest/asset['path']) as image:colors.append(image.getpixel((0,0)))
    assert colors == [(255,0,0),(0,0,255)]


@pytest.mark.parametrize('svg',['<svg xmlns="http://www.w3.org/2000/svg"><image href="https://example.invalid/x.png"/></svg>','<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>'])
def test_svg_active_resources_refused_even_without_vision(tmp_path,svg):
    source=tmp_path/'bad.svg';source.write_text(svg)
    p,dest=convert(tmp_path,source);failed(p,dest,source.stem)


def test_output_conflict_preserves_every_byte(tmp_path):
    source=tmp_path/'native.txt';source.write_text('Alpha 17 Beta 29')
    p,dest=convert(tmp_path,source);delivered(p,dest,source.stem)
    before={str(f.relative_to(dest)):f.read_bytes() for f in dest.rglob('*') if f.is_file()}
    p,_=convert(tmp_path,source,'--vision','on','--vision-backend','openai-compatible','--base-url','http://127.0.0.1:1/v1','--vision-model','unavailable')
    assert p.returncode == 2
    assert {str(f.relative_to(dest)):f.read_bytes() for f in dest.rglob('*') if f.is_file()} == before


def test_native_pdf_and_page_images_keep_page_ratio(tmp_path):
    from reportlab.pdfgen import canvas
    source=tmp_path/'digital.pdf';c=canvas.Canvas(str(source),pagesize=(300,600));c.drawString(20,550,'Alpha 17 Beta 29');c.showPage();c.save()
    p,dest=convert(tmp_path,source,'--page-images');md,report=delivered(p,dest,source.stem)
    assert 'Alpha' in md and 'Beta' in md
    assert report['engine'] == 'markitdown'
    assert len(report['assets']) == 1
    with Image.open(dest/report['assets'][0]['path']) as im: assert abs(im.height/im.width-2) < .01


def test_blank_pdf_can_succeed_without_invented_content(tmp_path):
    from reportlab.pdfgen import canvas
    source=tmp_path/'blank.pdf';c=canvas.Canvas(str(source));c.showPage();c.save()
    p,dest=convert(tmp_path,source);md,report=delivered(p,dest,source.stem)
    assert not report['assets']
    assert report['engine'] == 'markitdown'


@pytest.mark.parametrize('args',[['--page-images'],['--vision-timeout','0'],['--vision-timeout','nan']])
def test_invalid_conversion_options_leave_no_output(tmp_path,args):
    source=tmp_path/'native.txt';source.write_text('hello')
    p,dest=convert(tmp_path,source,*args)
    assert p.returncode == 2
    assert not dest.exists()


def test_markdown_images_preserve_source_tables_and_code(tmp_path):
    (tmp_path/'image.png').write_bytes(png('red'))
    source = tmp_path/'native.md'
    before = '# Heading\n\n| Name | Value |\n|---|---|\n| Alpha | 17 |\n\n'
    after = '\n\n```md\n![not an image](https://example.invalid/no.png)\n```\n'
    source.write_text(before + '![real](image.png)' + after)
    p,dest=convert(tmp_path,source);md,report=delivered(p,dest,source.stem)
    assert md.startswith(before) and md.endswith(after)
    assert len(report['assets']) == 1


def test_invalid_embedded_image_never_falls_back_to_placeholder(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text='Before bad image'
    slide.shapes.add_picture(BytesIO(png('red')), Inches(1), Inches(1))
    buffer=BytesIO();prs.save(buffer)
    source=tmp_path/'bad.pptx'
    with zipfile.ZipFile(buffer) as original, zipfile.ZipFile(source,'w') as broken:
        for member in original.infolist():
            broken.writestr(member, b'not an image' if member.filename.startswith('ppt/media/') else original.read(member))
    p,dest=convert(tmp_path,source);report=failed(p,dest,source.stem)
    assert 'slide:1/shape:' in p.stderr


def test_concurrent_output_claim_has_one_winner(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    source=tmp_path/'native.txt';source.write_text('Alpha 17 Beta 29')
    with ThreadPoolExecutor(max_workers=2) as pool:
        results=list(pool.map(lambda _: convert(tmp_path,source),range(2)))
    assert sorted(p.returncode for p,_ in results) == [0,2]
    winner,dest=next((p,d) for p,d in results if p.returncode == 0)
    md,report=delivered(winner,dest,source.stem)
    assert md == source.read_text()


def test_model_text_cannot_escape_into_executable_markdown(tmp_path, monkeypatch):
    sys.path.insert(0, str(RUNNER.parent))
    import vision
    from assets import AssetWriter
    from office_converters import AssetPptxConverter
    from markitdown import StreamInfo
    from pptx import Presentation
    from pptx.util import Inches
    from markdown_it import MarkdownIt
    from bs4 import BeautifulSoup
    transcription = '```\n<script>alert(1)</script>\n![upload](https://example.invalid/leak)\n```\n' + '完整转写' * 80
    monkeypatch.setattr(vision, 'analyze_image', lambda *args, **kwargs: vision.VisionResult(transcription, '<img src=\"https://example.invalid/leak\">', 'fixture'))
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.add_picture(BytesIO(png('red')),Inches(1),Inches(1))
    stream=BytesIO();prs.save(stream);stream.seek(0)
    writer=AssetWriter(tmp_path,vision=True)
    md=AssetPptxConverter(writer).convert(stream,StreamInfo(extension='.pptx')).markdown
    html=BeautifulSoup(MarkdownIt('commonmark',{'html':True}).render(md),'html.parser')
    assert not html.find_all('script')
    assert len(html.find_all('img')) == 1
    assert html.find('pre').get_text() == transcription


def test_visual_cache_retains_occurrences_but_changes_with_model(tmp_path, monkeypatch):
    sys.path.insert(0,str(RUNNER.parent))
    import vision
    from assets import AssetWriter
    calls=[]
    def analyze(path, **kwargs):
        calls.append(kwargs['model'])
        return vision.VisionResult('all text',kwargs['model'],kwargs['model'])
    monkeypatch.setattr(vision,'analyze_image',analyze)
    writer=AssetWriter(tmp_path,vision=True,model='first')
    data=png('red')
    first=writer.add_image(data,None,'slide:1')
    second=writer.add_image(data,None,'slide:2')
    writer.model='second'
    third=writer.add_image(data,None,'slide:3')
    assert calls == ['first','second']
    assert first == second and 'second' in third
    assert writer.entries[0]['source_refs'] == ['slide:1','slide:2','slide:3']
